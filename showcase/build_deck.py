"""Generate the Problem vs Solution deck from the same story as the web page.

Run: python3 showcase/build_deck.py
Output: showcase/problem-solution.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pathlib import Path

OUT = Path(__file__).with_name("problem-solution.pptx")

INK = RGBColor(0x16, 0x20, 0x2E)
MUTED = RGBColor(0x5B, 0x66, 0x76)
MUTED2 = RGBColor(0x8A, 0x94, 0xA4)
ACCENT = RGBColor(0x1D, 0x4E, 0xD8)
ACCENT_SOFT = RGBColor(0xE8, 0xEE, 0xFF)
BAD = RGBColor(0xC0, 0x39, 0x2B)
BAD_SOFT = RGBColor(0xFD, 0xEC, 0xEA)
GOOD = RGBColor(0x0F, 0x7A, 0x4D)
GOOD_SOFT = RGBColor(0xE2, 0xF5, 0xEC)
BORDER = RGBColor(0xE3, 0xE7, 0xEF)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
BG = RGBColor(0xF5, 0xF7, 0xFB)
FONT = "Segoe UI"
SLIDE_W = 13.333
SLIDE_H = 7.5

prs = Presentation()
prs.slide_width = Inches(SLIDE_W)
prs.slide_height = Inches(SLIDE_H)
BLANK = prs.slide_layouts[6]


def slide():
    s = prs.slides.add_slide(BLANK)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    bg.line.fill.background()
    bg.shadow.inherit = False
    return s


def box(s, x, y, w, h, fill=CARD, line=BORDER, radius=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    shp = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(0.75)
    shp.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    return shp


def text(s, x, y, w, h, items, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = it.get("align", align)
        if it.get("space_before"):
            p.space_before = Pt(it["space_before"])
        if it.get("space_after"):
            p.space_after = Pt(it["space_after"])
        run = p.add_run()
        run.text = it["t"]
        run.font.name = FONT
        run.font.size = Pt(it.get("s", 14))
        run.font.bold = it.get("b", False)
        run.font.color.rgb = it.get("c", INK)
    return tb


def shape_text(shp, items, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.05)
    tf.margin_bottom = Inches(0.05)
    first = True
    for it in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = it.get("align", align)
        p.space_after = Pt(it.get("space_after", 1))
        if it.get("space_before"):
            p.space_before = Pt(it["space_before"])
        run = p.add_run()
        run.text = it["t"]
        run.font.name = FONT
        run.font.size = Pt(it.get("s", 12))
        run.font.bold = it.get("b", False)
        run.font.color.rgb = it.get("c", INK)
    return shp


def down_arrow(s, x, y, w=0.3, h=0.22):
    a = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid()
    a.fill.fore_color.rgb = MUTED2
    a.line.fill.background()
    a.shadow.inherit = False
    return a


def right_arrow(s, x, y, w, h, color):
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    a.fill.solid()
    a.fill.fore_color.rgb = color
    a.line.fill.background()
    a.shadow.inherit = False
    return a


CURRENT_STEPS = [
    ("Business user", "Has a question but no way to query the data directly.", "Starts request", False),
    ("Queries the system", "Raises a ticket describing what is needed.", "2 - 4 h", True),
    ("OPS team queue", "Waits in the support backlog behind other tickets.", "1 - 2 days", True),
    ("OPS validates", "Scope, fields, and permissions are clarified over email.", "2 - 4 h", True),
    ("Transferred to Dev", "A developer picks it up when capacity allows.", "1 - 3 days", True),
    ("Response returned", "The result is pasted back and reaches the user.", "1 - 2 h", False),
]
CURRENT_WAITS = ["Queue + triage\n4 - 8 h", "Backlog wait\n1 - 2 d", "Handoff\n4 - 8 h", "Dev capacity\n1 - 3 d", "Follow-up\n2 - 6 h"]

SOLUTION_STEPS = [
    ("Business user", "Asks in plain English, no query language needed.", "Under 1 min", False),
    ("AI chat system", "Understands intent locally: cache, slots, routing.", "Cache < 10 ms", False),
    ("Business filter", "Applies the schema, sensitivity flags, guardrails.", "Guardrail < 1 ms", False),
    ("Output shared", "Governed answer, table, briefing, or export.", "1.5 - 3 s if AI used", False),
]
SOLUTION_WAITS = ["Typed, no ticket\nInstant", "In-memory\nInstant", "Validated\nInstant"]


def header(s, kicker, title, total=None, total_kind=None):
    text(s, 0.55, 0.34, 8.6, 0.3, [{"t": kicker, "s": 10, "b": True, "c": ACCENT}])
    text(s, 0.55, 0.6, 9.2, 0.6, [{"t": title, "s": 25, "b": True, "c": INK}])
    if total:
        fill = BAD_SOFT if total_kind == "bad" else GOOD_SOFT
        ink = BAD if total_kind == "bad" else GOOD
        b = box(s, 10.6, 0.42, 2.25, 0.92, fill=fill, line=None)
        shape_text(b, [
            {"t": "TYPICAL TURNAROUND", "s": 8.5, "b": True, "c": ink, "align": PP_ALIGN.CENTER},
            {"t": total, "s": 21, "b": True, "c": ink, "align": PP_ALIGN.CENTER},
        ], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def flow_slide(kicker, title, steps, waits, total, kind):
    s = slide()
    header(s, kicker, title, total, kind)
    body_ink = BAD if kind == "bad" else ACCENT
    arrow_color = RGBColor(0xE3, 0xA4, 0x9C) if kind == "bad" else RGBColor(0x9F, 0xD9, 0xC0)

    top, h = 1.85, 2.75
    if len(steps) == 6:
        node_w, gap = 1.72, 0.40
    else:
        node_w, gap = 2.45, 0.90
    x = 0.5
    centers = []
    for i, (t, body, dur, queue) in enumerate(steps):
        shp = box(s, x, top, node_w, h)
        if queue:
            shp.line.color.rgb = arrow_color
        items = []
        if kind == "bad":
            items.append({"t": "WAITING" if queue else " ", "s": 7.5, "b": True, "c": BAD if queue else BG, "space_after": 3})
        tag_fill = GOOD_SOFT if dur.lower().startswith(("cache", "guardrail")) or "1.5" in dur else (ACCENT_SOFT if kind == "good" else BAD_SOFT)
        items += [
            {"t": t, "s": 11, "b": True, "c": INK, "space_after": 3},
            {"t": body, "s": 8.5, "c": MUTED, "space_after": 5},
        ]
        shape_text(shp, items, anchor=MSO_ANCHOR.TOP)
        chip = box(s, x + 0.12, top + h - 0.5, node_w - 0.24, 0.34, fill=tag_fill, line=None, radius=0.5)
        shape_text(chip, [{"t": dur, "s": 8.5, "b": True, "c": body_ink, "align": PP_ALIGN.CENTER}], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        centers.append(x + node_w / 2)
        x += node_w
        if i < len(steps) - 1:
            right_arrow(s, x + (gap - 0.34) / 2, top + h / 2 - 0.11, 0.34, 0.22, arrow_color)
            text(s, x + gap / 2 - 0.72, top + h / 2 + 0.16, 1.44, 0.6,
                 [{"t": waits[i], "s": 8, "b": True, "c": body_ink, "align": PP_ALIGN.CENTER}], align=PP_ALIGN.CENTER)
            x += gap
    return s


# 1. Title
s = slide()
text(s, 0.8, 1.5, 8.0, 0.4, [{"t": "MULTI-AGENT MONGODB NLP", "s": 12, "b": True, "c": ACCENT}])
text(s, 0.8, 1.95, 11.7, 1.9, [
    {"t": "From days of handoffs", "s": 40, "b": True, "c": BAD, "space_after": 2},
    {"t": "to seconds of self-service", "s": 40, "b": True, "c": GOOD},
])
text(s, 0.8, 4.15, 10.5, 0.9, [{
    "t": "The same business question takes two very different paths. One moves through people and queues. "
         "The other is resolved inside a governed AI workspace. The difference is time, cost, and control.",
    "s": 15, "c": MUTED}])
for i, (label, val, fill, ink) in enumerate([
    ("TODAY", "2 - 5 days", BAD_SOFT, BAD),
    ("WITH THE PLATFORM", "Seconds", GOOD_SOFT, GOOD),
]):
    x = 0.8 + i * 3.3
    b = box(s, x, 5.35, 3.0, 1.0, fill=fill, line=None)
    shape_text(b, [
        {"t": label, "s": 9, "b": True, "c": ink},
        {"t": val, "s": 20, "b": True, "c": ink},
    ], anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.8, 6.65, 11.5, 0.4, [{
    "t": "Illustrative request lifecycle, not a measured production baseline. The platform path assumes no sensitive-data approval pause.",
    "s": 9, "c": MUTED2}])

# 2. Comparison
s = slide()
header(s, "AT A GLANCE", "Two paths for the same question")
rows = [("Today", "People and queues", 100, BAD, "2 - 5 days"), ("With the platform", "Governed AI workspace", 4, GOOD, "Seconds")]
y = 2.1
for label, sub, pct, color, amt in rows:
    text(s, 0.7, y, 2.6, 0.5, [{"t": label, "s": 14, "b": True, "c": INK}, {"t": sub, "s": 10, "c": MUTED2}])
    track = box(s, 3.35, y + 0.06, 7.6, 0.42, fill=RGBColor(0xEE, 0xF1, 0xF7), line=None, radius=0.5)
    bar = box(s, 3.35, y + 0.06, max(0.18, 7.6 * pct / 100.0), 0.42, fill=color, line=None, radius=0.5)
    text(s, 11.15, y, 1.6, 0.5, [{"t": amt, "s": 16, "b": True, "c": color}], align=PP_ALIGN.RIGHT)
    y += 1.5
text(s, 0.7, 5.35, 11.9, 0.9, [{
    "t": "The current path is slow because the answer depends on people being free. The platform removes that dependency: "
         "simple questions are answered locally and instantly, hard ones are reasoned and checked, and sensitive requests stop for a human by design.",
    "s": 13, "c": MUTED}])

# 3. Current
flow_slide("CURRENT PROCESS", "The long road: every question is a ticket", CURRENT_STEPS, CURRENT_WAITS, "2 - 5 days", "bad")

# 4. Solution
flow_slide("WITH THE PLATFORM", "The direct path: one governed workspace", SOLUTION_STEPS, SOLUTION_WAITS, "Seconds", "good")

# 5. Architecture
s = slide()
header(s, "SYSTEM DESIGN", "The system behind the direct path")
layers = [
    ("Experience layer", "Angular 21 SPA", RGBColor(0x6D, 0x8C, 0xFF),
     ["Personal greeting", "Quick-action chips", "Agent activity stream", "Results grid", "Approval portal"]),
    ("Intelligence layer", "Hybrid NLP and agents", ACCENT,
     ["Semantic cache, bge-small", "Slot extraction", "Scriban templates", "NVIDIA NIM for complex queries", "AST validation and self-correction"]),
    ("Governance layer", "Controls, not paperwork", RGBColor(0xB7, 0x79, 0x1F),
     ["Field sensitivity registry", "Read-only guardrails", "Approval and override", "Append-only audit trail"]),
    ("Data layer", "MongoDB and enterprise APIs", GOOD,
     ["sample_airbnb listings", "Enterprise Core REST API", "Access requests", "Audit logs"]),
]
y = 1.45
inner_left, inner_right = 3.7, 12.6
for i, (name, sub, color, comps) in enumerate(layers):
    rows, cur, cur_w = [], [], 0.0
    for comp in comps:
        cw = 0.20 + 0.10 * len(comp)
        if cur and cur_w + 0.12 + cw > (inner_right - inner_left):
            rows.append(cur)
            cur, cur_w = [], 0.0
        cur.append((comp, cw))
        cur_w += (0.12 if cur_w else 0.0) + cw
    if cur:
        rows.append(cur)
    band_h = 0.18 + len(rows) * 0.42 + (len(rows) - 1) * 0.12 + 0.18
    box(s, 0.6, y, 12.1, band_h, fill=CARD)
    box(s, 0.6, y, 0.09, band_h, fill=color, line=None, shape=MSO_SHAPE.RECTANGLE)
    text(s, 0.85, y + 0.18, 2.6, band_h - 0.3, [{"t": name, "s": 13, "b": True, "c": INK, "space_after": 2}, {"t": sub, "s": 10, "c": MUTED2}])
    ry = y + 0.18
    for row in rows:
        cx = inner_left
        for comp, cw in row:
            c = box(s, cx, ry, cw, 0.42, fill=RGBColor(0xF9, 0xFA, 0xFC), radius=0.3)
            shape_text(c, [{"t": comp, "s": 9.5, "b": True, "c": MUTED, "align": PP_ALIGN.CENTER}], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            cx += cw + 0.12
        ry += 0.54
    y += band_h
    if i < len(layers) - 1:
        down_arrow(s, 6.5, y, 0.3, 0.14)
        y += 0.14

# 6. Closing
s = slide()
header(s, "THE POINT", "Speed and governance stop being a trade-off")
points = [
    ("Simple questions", "Answered locally and instantly, at zero model cost."),
    ("Hard questions", "Reasoned by the model, validated, and self-corrected up to three times."),
    ("Sensitive requests", "Paused for an approver by design, with a manager override path."),
    ("Every action", "Recorded in an append-only audit trail with actor and justification."),
]
y = 1.75
for t, body in points:
    dot = box(s, 0.75, y + 0.1, 0.16, 0.16, fill=ACCENT, line=None, shape=MSO_SHAPE.OVAL)
    text(s, 1.1, y, 11.3, 0.7, [{"t": t, "s": 15, "b": True, "c": INK, "space_after": 2}, {"t": body, "s": 12, "c": MUTED}])
    y += 0.95
text(s, 0.75, 6.55, 11.9, 0.5, [{
    "t": "Design and process illustration. Times are illustrative of a typical enterprise request lifecycle, not measured production results.",
    "s": 9, "c": MUTED2}])

prs.save(str(OUT))
print("wrote", OUT)

limits_w = prs.slide_width
limits_h = prs.slide_height
violations = []
for idx, sl in enumerate(prs.slides, start=1):
    for shp in sl.shapes:
        if shp.left < 0 or shp.top < 0 or shp.left + shp.width > limits_w + Emu(10000) or shp.top + shp.height > limits_h + Emu(10000):
            violations.append((idx, shp.shape_type, shp.left, shp.top, shp.width, shp.height))
if violations:
    print("BOUNDS VIOLATIONS:", len(violations))
    for v in violations:
        print("  slide", v[0], "left", round(v[2] / 914400, 2), "top", round(v[3] / 914400, 2), "right", round((v[2] + v[4]) / 914400, 2), "bottom", round((v[3] + v[5]) / 914400, 2))
    raise SystemExit(1)
print("bounds check passed across", len(prs.slides._sldIdLst), "slides")
